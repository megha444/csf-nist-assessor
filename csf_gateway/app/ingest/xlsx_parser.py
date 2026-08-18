"""
XLSX → markdown via openpyxl.

XLSX is authored (not scanned), so openpyxl handles it well without needing a
multimodal LLM. Each worksheet becomes a heading + one markdown table, in
workbook sheet order — mirrors docx_parser.py's table-to-markdown contract so
the same chunker/tagger/embedder pipeline downstream doesn't need to know the
source format.
"""
from __future__ import annotations

import logging
from io import BytesIO

from openpyxl import load_workbook

log = logging.getLogger("csf-gateway.ingest.xlsx")


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("|", "\\|").replace("\n", " ")


def _sheet_to_md(rows: list[tuple]) -> str:
    rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
    if not rows:
        return ""
    n_cols = max(len(r) for r in rows)
    lines: list[str] = []
    for i, row in enumerate(rows):
        cells = [_cell_to_str(row[j]) if j < len(row) else "" for j in range(n_cols)]
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("| " + " | ".join(["---"] * n_cols) + " |")
    return "\n".join(lines)


def xlsx_to_markdown(xlsx_bytes: bytes) -> str:
    """Convert an .xlsx file (bytes) to markdown. One `##` heading + table per non-empty sheet."""
    wb = load_workbook(BytesIO(xlsx_bytes), data_only=True, read_only=True)
    out: list[str] = []
    sheet_count = 0
    for ws in wb.worksheets:
        table_md = _sheet_to_md(list(ws.iter_rows(values_only=True)))
        if not table_md:
            continue
        sheet_count += 1
        out.append(f"## {ws.title}")
        out.append("")
        out.append(table_md)
        out.append("")
    wb.close()
    text = "\n".join(out).strip()
    log.info("xlsx→md: %d bytes, %d non-empty sheets → %d chars", len(xlsx_bytes), sheet_count, len(text))
    return text
