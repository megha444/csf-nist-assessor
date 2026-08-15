"""
Extract per-Function content from Sakina's `Ref Docs/*.xlsx` files into YAML.

Reads:
  Ref Docs/NIST Rubrics.xlsx                      sheet <Function> (e.g. "Govern")
  Ref Docs/CSF 2.0-Implementation_Examples.xlsx   sheet "Sheet1", filtered to <FUNCTION>

Writes:
  csf_specialists/rubrics/<function>.yaml
  csf_specialists/implementation_examples/<function>.yaml

Usage:
  poetry run python infra/build_rubric_yaml.py --function govern
  poetry run python infra/build_rubric_yaml.py --function identify
  poetry run python infra/build_rubric_yaml.py --function all       # do every ready Function

Idempotent — safe to re-run whenever Sakina updates the source spreadsheets.
Missing rubric sheet for a Function → prints a clear "waiting on Sakina" note.
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
REF_DIR = ROOT / "Ref Docs"
RUBRICS_XLSX = REF_DIR / "NIST Rubrics.xlsx"
EXAMPLES_XLSX = REF_DIR / "CSF 2.0-Implementation_Examples.xlsx"

OUT_RUBRIC_DIR = ROOT / "csf_specialists" / "rubrics"
OUT_EXAMPLES_DIR = ROOT / "csf_specialists" / "implementation_examples"

DEFAULT_CATEGORY_IMPORTANCE = 2


@dataclass(frozen=True)
class FunctionSpec:
    key: str                # our lowercase name, e.g. "govern"
    sheet: str              # NIST Rubrics.xlsx sheet name, e.g. "Govern"
    code_prefix: str        # CSF code prefix, e.g. "GV"
    examples_label: str     # what the Implementation Examples file's Function column starts with


FUNCTIONS: dict[str, FunctionSpec] = {
    "govern":   FunctionSpec("govern",   "Govern",   "GV", "GOVERN"),
    "identify": FunctionSpec("identify", "Identify", "ID", "IDENTIFY"),
    "protect":  FunctionSpec("protect",  "Protect",  "PR", "PROTECT"),
    "detect":   FunctionSpec("detect",   "Detect",   "DE", "DETECT"),
    "respond":  FunctionSpec("respond",  "Respond",  "RS", "RESPOND"),
    "recover":  FunctionSpec("recover",  "Recover",  "RC", "RECOVER"),
}


def extract_rubric(spec: FunctionSpec) -> dict | None:
    """
    Read the Function's rubric sheet. Returns None if the sheet doesn't exist
    yet (e.g. Sakina hasn't authored it) — caller decides what to log.
    """
    wb = openpyxl.load_workbook(RUBRICS_XLSX, data_only=True)
    if spec.sheet not in wb.sheetnames:
        return None
    ws = wb[spec.sheet]

    cat_header_re = re.compile(rf"^\s*({spec.code_prefix}\.[A-Z]{{2}})\s*-\s*(.+?)\s*$")

    categories: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        header, *level_cells = row[:6]
        if not header:
            continue
        m = cat_header_re.match(str(header))
        if not m:
            continue
        code, name = m.group(1), m.group(2)
        levels: dict[int, str] = {}
        for i, cell in enumerate(level_cells, start=1):
            if not cell:
                continue
            levels[i] = str(cell).strip()
        categories[code] = {
            "name": name,
            "importance": DEFAULT_CATEGORY_IMPORTANCE,
            "levels": levels,
        }

    if not categories:
        return None  # sheet present but empty / only header — treat as not yet delivered
    return {"function": spec.key, "categories": categories}


def extract_examples(spec: FunctionSpec) -> dict | None:
    """
    Walk Sheet1, filter to this Function's rows, accumulate examples per Subcategory.
    Returns None if the source file is missing.
    """
    if not EXAMPLES_XLSX.exists():
        return None
    wb = openpyxl.load_workbook(EXAMPLES_XLSX, data_only=True)
    ws = wb["Sheet1"]

    sub_id_re = re.compile(rf"^\s*({spec.code_prefix}\.[A-Z]{{2}}-\d{{1,2}})\s*:\s*(.+?)\s*$")

    subcategories: dict[str, dict] = {}
    in_this_function = False
    current_sub: str | None = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        function_cell, category_cell, sub_cell, example_cell = row[:4]

        if function_cell:
            s = str(function_cell)
            in_this_function = s.startswith(spec.examples_label)

        if not in_this_function:
            continue

        if sub_cell:
            m = sub_id_re.match(str(sub_cell))
            if m:
                sub_id, sub_name = m.group(1), m.group(2)
                current_sub = sub_id
                subcategories.setdefault(sub_id, {"name": sub_name, "examples": []})

        if example_cell and current_sub:
            text = str(example_cell).replace("\xa0", " ").strip()
            if text:
                subcategories[current_sub]["examples"].append(text)

    return dict(sorted(subcategories.items())) if subcategories else None


def write_yaml(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w") as f:
        yaml.safe_dump(
            data, f, sort_keys=False, allow_unicode=True, width=100, default_flow_style=False,
        )
    print(f"    wrote {path.relative_to(ROOT)}  ({path.stat().st_size} bytes)")


def process(spec: FunctionSpec) -> tuple[bool, bool]:
    """Returns (rubric_written, examples_written)."""
    print(f"\n=== {spec.key} ===")

    rubric = extract_rubric(spec)
    if rubric is None:
        print(f"  [rubric]   PENDING from Sakina — sheet {spec.sheet!r} not found or empty")
        rubric_written = False
    else:
        cats = rubric["categories"]
        print(f"  [rubric]   {len(cats)} Categories: {list(cats)}")
        for code, cat in cats.items():
            if set(cat["levels"]) != {1, 2, 3, 4, 5}:
                print(f"    ⚠  {code} has levels {sorted(cat['levels'])} (expected 1-5)")
        write_yaml(OUT_RUBRIC_DIR / f"{spec.key}.yaml", rubric)
        rubric_written = True

    examples = extract_examples(spec)
    if examples is None:
        print(f"  [examples] no {spec.examples_label} rows found in Implementation Examples file")
        examples_written = False
    else:
        total = sum(len(v["examples"]) for v in examples.values())
        print(f"  [examples] {len(examples)} Subcategories, {total} examples")
        for sub_id, entry in examples.items():
            if not entry["examples"]:
                print(f"    ⚠  {sub_id} has 0 examples")
        write_yaml(OUT_EXAMPLES_DIR / f"{spec.key}.yaml", examples)
        examples_written = True

    return rubric_written, examples_written


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--function",
        required=True,
        choices=list(FUNCTIONS) + ["all"],
        help="Function to extract (or 'all' to try every one)",
    )
    args = parser.parse_args()

    if not RUBRICS_XLSX.exists():
        print(f"[error] {RUBRICS_XLSX} not found", file=sys.stderr)
        return 1

    targets = list(FUNCTIONS.values()) if args.function == "all" else [FUNCTIONS[args.function]]

    summary: list[tuple[str, bool, bool]] = []
    for spec in targets:
        r, e = process(spec)
        summary.append((spec.key, r, e))

    print("\n=== summary ===")
    for key, r, e in summary:
        rmark = "✅" if r else "⏳"
        emark = "✅" if e else "⏳"
        print(f"  {key:<9}  rubric {rmark}    examples {emark}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
